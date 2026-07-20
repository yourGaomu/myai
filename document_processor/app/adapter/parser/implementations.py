from __future__ import annotations

import re
from dataclasses import dataclass
from html.parser import HTMLParser

from ...errors import ProcessingError
from ...models import DocumentSource, ParsedDocument, ParsingProfile, Section


def _sections_from_markdown(text: str) -> list[Section]:
    headings: list[tuple[int, int, str]] = []
    char_offset = 0
    byte_offset = 0
    for line in text.splitlines(keepends=True):
        match = re.match(r"^\s{0,3}#{1,6}\s+(.+?)\s*$", line.rstrip("\r\n"))
        if match:
            headings.append((char_offset, byte_offset, match.group(1).strip()))
        char_offset += len(line)
        byte_offset += len(line.encode("utf-8"))

    if not headings:
        return [Section("", text, 0, len(text.encode("utf-8")), 0)] if text else []
    sections: list[Section] = []
    for index, (char_start, byte_start, heading) in enumerate(headings):
        char_end = headings[index + 1][0] if index + 1 < len(headings) else len(text)
        byte_end = headings[index + 1][1] if index + 1 < len(headings) else len(text.encode("utf-8"))
        sections.append(Section(heading, text[char_start:char_end], byte_start, byte_end, 0))
    if headings[0][0] > 0:
        sections.insert(0, Section("", text[: headings[0][0]], 0, headings[0][1], 0))
    return sections


@dataclass(frozen=True)
class TextParser:
    parser_id: str
    version: str
    content_types: frozenset[str]
    markdown: bool = False

    def supports(self, content_type: str) -> bool:
        return content_type.lower().split(";", 1)[0].strip() in self.content_types

    def parse(self, source: DocumentSource, profile: ParsingProfile) -> ParsedDocument:
        text = source.path.read_text(encoding="utf-8-sig", errors="replace")
        sections = _sections_from_markdown(text) if self.markdown else ([Section("", text, 0, len(text.encode("utf-8")), 0)] if text else [])
        return ParsedDocument(text=text, sections=sections)


class _HTMLTextParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.heading_parts: list[str] = []
        self.headings: list[str] = []
        self.in_heading = False

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            self.heading_parts = []
            self.in_heading = True
        elif tag in {"p", "div", "li", "br", "section", "article"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"h1", "h2", "h3", "h4", "h5", "h6"} and self.in_heading:
            heading = " ".join("".join(self.heading_parts).split())
            if heading:
                self.headings.append(heading)
                self.parts.append(heading)
                self.parts.append("\n")
            self.heading_parts = []
            self.in_heading = False

    def handle_data(self, data: str) -> None:
        if self.in_heading:
            self.heading_parts.append(data)
        else:
            self.parts.append(data)

    def document(self) -> ParsedDocument:
        text = re.sub(r"[ \t]+\n", "\n", "".join(self.parts))
        text = re.sub(r"\n{3,}", "\n\n", text).strip()
        heading = self.headings[0] if self.headings else ""
        sections = [Section(heading, text, 0, len(text.encode("utf-8")), 0)] if text else []
        return ParsedDocument(text=text, sections=sections)


@dataclass(frozen=True)
class HTMLParserAdapter:
    parser_id: str = "python-html"
    version: str = "1"

    def supports(self, content_type: str) -> bool:
        return content_type.lower().split(";", 1)[0].strip() in {"text/html", "application/xhtml+xml"}

    def parse(self, source: DocumentSource, profile: ParsingProfile) -> ParsedDocument:
        parser = _HTMLTextParser()
        parser.feed(source.path.read_text(encoding="utf-8-sig", errors="replace"))
        parser.close()
        return parser.document()


@dataclass(frozen=True)
class PDFParser:
    parser_id: str = "python-pymupdf"
    version: str = "1"

    def supports(self, content_type: str) -> bool:
        return content_type.lower().split(";", 1)[0].strip() == "application/pdf"

    def parse(self, source: DocumentSource, profile: ParsingProfile) -> ParsedDocument:
        try:
            import fitz
        except ImportError as exc:
            raise ProcessingError("PDF parser requires optional PyMuPDF dependency") from exc
        pdf = fitz.open(filename=str(source.path), filetype="pdf")
        pages: list[str] = []
        sections: list[Section] = []
        offset = 0
        for page_number, page in enumerate(pdf, start=1):
            page_text = page.get_text("text")
            if pages:
                offset += 1
            pages.append(page_text)
            sections.append(Section("", page_text, offset, offset + len(page_text), page_number))
            offset += len(page_text.encode("utf-8"))
        pdf.close()
        return ParsedDocument("\n".join(pages), sections)


@dataclass(frozen=True)
class DOCXParser:
    parser_id: str = "python-docx"
    version: str = "1"

    def supports(self, content_type: str) -> bool:
        return content_type.lower().split(";", 1)[0].strip() in {
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/msword",
        }

    def parse(self, source: DocumentSource, profile: ParsingProfile) -> ParsedDocument:
        try:
            from docx import Document
        except ImportError as exc:
            raise ProcessingError("DOCX parser requires optional python-docx dependency") from exc
        document = Document(str(source.path))
        lines: list[str] = []
        sections: list[Section] = []
        offset = 0
        for paragraph in document.paragraphs:
            value = paragraph.text.strip()
            if not value:
                continue
            if lines:
                offset += 1
            lines.append(value)
            heading = value if paragraph.style.name.lower().startswith("heading") else ""
            value_size = len(value.encode("utf-8"))
            sections.append(Section(heading, value, offset, offset + value_size, 0))
            offset += value_size
        return ParsedDocument("\n".join(lines), sections)
